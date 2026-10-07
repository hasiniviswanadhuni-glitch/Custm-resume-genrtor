import os
import uuid
from typing import Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.app.config import settings

class ExportService:
    """
    High-fidelity ATS-friendly Document Exporter.
    Generates PDF (ReportLab) and DOCX (python-docx) files from TailoredResumeContent.
    Validates file creation, integrity, and byte size.
    """

    def __init__(self):
        os.makedirs(settings.STORAGE_DIR, exist_ok=True)

    def export_to_pdf(self, content: Dict[str, Any], output_filename: str = None) -> str:
        if not output_filename:
            output_filename = f"resume_{uuid.uuid4().hex[:8]}.pdf"
        filepath = os.path.join(settings.STORAGE_DIR, output_filename)

        doc = SimpleDocTemplate(
            filepath,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Custom Typography
        title_style = ParagraphStyle(
            'ResumeTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1E293B'),
            alignment=1 # Center
        )
        contact_style = ParagraphStyle(
            'ResumeContact',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#64748B'),
            alignment=1 # Center
        )
        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=8,
            spaceAfter=3
        )
        body_style = ParagraphStyle(
            'ResumeBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#334155')
        )
        bullet_style = ParagraphStyle(
            'ResumeBullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            leftIndent=12,
            textColor=colors.HexColor('#334155')
        )
        item_title_style = ParagraphStyle(
            'ItemTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=13,
            textColor=colors.HexColor('#1E293B')
        )

        story = []

        # 1. Header: Name & Contact Details
        full_name = content.get("full_name", "Candidate")
        story.append(Paragraph(full_name, title_style))
        story.append(Spacer(1, 4))

        contact_parts = []
        if content.get("email"): contact_parts.append(content["email"])
        if content.get("phone"): contact_parts.append(content["phone"])
        if content.get("location"): contact_parts.append(content["location"])
        if content.get("linkedin_url"): contact_parts.append(content["linkedin_url"])
        if content.get("github_url"): contact_parts.append(content["github_url"])

        contact_line = " | ".join(contact_parts)
        story.append(Paragraph(contact_line, contact_style))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceBefore=2, spaceAfter=8))

        # 2. Professional Summary
        summary = content.get("professional_summary")
        if summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceBefore=1, spaceAfter=4))
            story.append(Paragraph(summary, body_style))
            story.append(Spacer(1, 6))

        # 3. Skills
        skills = content.get("skills", [])
        if skills:
            story.append(Paragraph("CORE SKILLS & TECHNOLOGIES", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceBefore=1, spaceAfter=4))
            skill_names = [s.get("name", "") if isinstance(s, dict) else str(s) for s in skills]
            story.append(Paragraph(", ".join(skill_names), body_style))
            story.append(Spacer(1, 6))

        # 4. Experience
        experience = content.get("experience", [])
        if experience:
            story.append(Paragraph("WORK EXPERIENCE", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceBefore=1, spaceAfter=4))
            for exp in experience:
                pos = exp.get("position", "")
                comp = exp.get("company", "")
                dates = f"{exp.get('start_date', '')} - {exp.get('end_date', 'Present')}"
                story.append(Paragraph(f"<b>{pos}</b> — {comp} <font color='#64748B' size=8>({dates})</font>", item_title_style))
                for bullet in exp.get("bullets", []):
                    story.append(Paragraph(f"• {bullet}", bullet_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 4))

        # 5. Projects
        projects = content.get("projects", [])
        if projects:
            story.append(Paragraph("TECHNICAL PROJECTS", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceBefore=1, spaceAfter=4))
            for proj in projects:
                p_title = proj.get("title", "")
                techs = ", ".join(proj.get("technologies", []))
                tech_line = f" [{techs}]" if techs else ""
                story.append(Paragraph(f"<b>{p_title}</b>{tech_line}", item_title_style))
                for bullet in proj.get("bullets", []):
                    story.append(Paragraph(f"• {bullet}", bullet_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 4))

        # 6. Education
        education = content.get("education", [])
        if education:
            story.append(Paragraph("EDUCATION", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceBefore=1, spaceAfter=4))
            for edu in education:
                deg = edu.get("degree", "")
                inst = edu.get("institution", "")
                dates = f"{edu.get('start_date', '')} - {edu.get('end_date', '')}"
                story.append(Paragraph(f"<b>{deg}</b> — {inst} <font color='#64748B' size=8>({dates})</font>", item_title_style))
            story.append(Spacer(1, 4))

        # 7. Certifications
        certs = content.get("certifications", [])
        if certs:
            story.append(Paragraph("CERTIFICATIONS", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#E2E8F0'), spaceBefore=1, spaceAfter=4))
            for c in certs:
                story.append(Paragraph(f"• <b>{c.get('name', '')}</b> ({c.get('issuer', '')})", bullet_style))

        doc.build(story)

        # Validate file existence and non-zero byte size
        if not os.path.exists(filepath) or os.path.getsize(filepath) == 0:
            raise RuntimeError(f"PDF export validation failed for: {filepath}")

        return filepath

    def export_to_docx(self, content: Dict[str, Any], output_filename: str = None) -> str:
        if not output_filename:
            output_filename = f"resume_{uuid.uuid4().hex[:8]}.docx"
        filepath = os.path.join(settings.STORAGE_DIR, output_filename)

        doc = Document()

        # Set 0.5-inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # Title
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_name = p_name.add_run(content.get("full_name", "Candidate"))
        run_name.font.size = Pt(18)
        run_name.font.bold = True
        run_name.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # Contact Details
        contact_parts = []
        if content.get("email"): contact_parts.append(content["email"])
        if content.get("phone"): contact_parts.append(content["phone"])
        if content.get("location"): contact_parts.append(content["location"])
        if content.get("linkedin_url"): contact_parts.append(content["linkedin_url"])
        if content.get("github_url"): contact_parts.append(content["github_url"])

        p_contact = doc.add_paragraph(" | ".join(contact_parts))
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_contact.paragraph_format.space_after = Pt(10)

        # Helper for Section Heading
        def add_heading(text: str):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            r = h.add_run(text.upper())
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        # Summary
        if content.get("professional_summary"):
            add_heading("Professional Summary")
            p_sum = doc.add_paragraph(content["professional_summary"])
            p_sum.paragraph_format.space_after = Pt(6)

        # Skills
        if content.get("skills"):
            add_heading("Core Skills & Technologies")
            skill_names = [s.get("name", "") if isinstance(s, dict) else str(s) for s in content["skills"]]
            p_sk = doc.add_paragraph(", ".join(skill_names))
            p_sk.paragraph_format.space_after = Pt(6)

        # Experience
        if content.get("experience"):
            add_heading("Work Experience")
            for exp in content["experience"]:
                p_exp = doc.add_paragraph()
                r_pos = p_exp.add_run(f"{exp.get('position', '')} — {exp.get('company', '')} ")
                r_pos.font.bold = True
                p_exp.add_run(f"({exp.get('start_date', '')} - {exp.get('end_date', 'Present')})")
                for b in exp.get("bullets", []):
                    doc.add_paragraph(b, style='List Bullet')

        # Projects
        if content.get("projects"):
            add_heading("Technical Projects")
            for proj in content["projects"]:
                p_proj = doc.add_paragraph()
                r_title = p_proj.add_run(proj.get("title", ""))
                r_title.font.bold = True
                techs = ", ".join(proj.get("technologies", []))
                if techs:
                    p_proj.add_run(f" [{techs}]")
                for b in proj.get("bullets", []):
                    doc.add_paragraph(b, style='List Bullet')

        # Education
        if content.get("education"):
            add_heading("Education")
            for edu in content["education"]:
                p_edu = doc.add_paragraph()
                r_deg = p_edu.add_run(f"{edu.get('degree', '')} — {edu.get('institution', '')} ")
                r_deg.font.bold = True
                p_edu.add_run(f"({edu.get('start_date', '')} - {edu.get('end_date', '')})")

        # Certifications
        if content.get("certifications"):
            add_heading("Certifications")
            for c in content["certifications"]:
                doc.add_paragraph(f"{c.get('name', '')} ({c.get('issuer', '')})", style='List Bullet')

        doc.save(filepath)

        # Validate file existence and non-zero byte size
        if not os.path.exists(filepath) or os.path.getsize(filepath) == 0:
            raise RuntimeError(f"DOCX export validation failed for: {filepath}")

        return filepath

export_service = ExportService()
